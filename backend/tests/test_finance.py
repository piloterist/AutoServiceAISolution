"""Tests for Настройки -> Финансы (cost items) and Финансы -> Расходы -
services/finance_service.py and its API."""

import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models.cost_expense import CostExpense
from app.models.cost_item import CostItem
from app.models.department import Department
from app.models.workshop import Workshop
from app.services import finance_import_service, finance_service
from app.services.finance_service import FinanceConflict, FinanceInvalid, FinanceNotFound

ITEMS_URL = "/api/v1/finance/cost-items"
EXPENSES_URL = "/api/v1/finance/expenses"


def _workshop(db: Session, dept_name: str = "Каховка", kind: str = "Кузовной") -> Workshop:
    dept = db.query(Department).filter_by(name=dept_name).first()
    if dept is None:
        dept = Department(name=dept_name)
        db.add(dept)
        db.flush()
    workshop = Workshop(
        department_id=dept.id,
        workshop_type=kind,
        posts_count=1,
        start_time="08:00:00",
        end_time="20:00:00",
        working_days=[0],
    )
    db.add(workshop)
    db.flush()
    return workshop


def test_cost_item_names_are_unique_and_deletable_only_while_unused(db_session: Session) -> None:
    item = finance_service.create_cost_item(db_session, name="Аренда", is_active=True)
    with pytest.raises(FinanceConflict):
        finance_service.create_cost_item(db_session, name="Аренда", is_active=True)

    other = finance_service.create_cost_item(db_session, name="Свет", is_active=True)
    finance_service.delete_cost_item(db_session, other.id)  # unused -> fine
    with pytest.raises(FinanceNotFound):
        finance_service.delete_cost_item(db_session, other.id)

    workshop = _workshop(db_session)
    finance_service.set_expense(
        db_session,
        cost_item_id=item.id,
        workshop_id=workshop.id,
        year=2026,
        month=9,
        amount=Decimal("100"),
    )
    with pytest.raises(FinanceConflict):
        finance_service.delete_cost_item(db_session, item.id)


def test_amounts_only_for_active_items_zero_clears_and_cells_upsert(db_session: Session) -> None:
    workshop = _workshop(db_session)
    active = finance_service.create_cost_item(db_session, name="Аренда", is_active=True)
    inactive = finance_service.create_cost_item(db_session, name="Старая", is_active=False)

    def put(item: CostItem, amount: str) -> None:
        finance_service.set_expense(
            db_session,
            cost_item_id=item.id,
            workshop_id=workshop.id,
            year=2026,
            month=9,
            amount=Decimal(amount),
        )

    with pytest.raises(FinanceInvalid):
        put(inactive, "5")

    put(active, "100")
    put(active, "250.50")  # same cell: updated in place
    assert db_session.query(CostExpense).count() == 1
    assert finance_service.monthly_expense_totals(db_session, 2026) == {
        (workshop.id, 9): Decimal("250.50")
    }

    put(active, "0")  # zero clears the cell
    assert db_session.query(CostExpense).count() == 0


def test_year_view_has_one_block_per_workshop_and_hides_unused_inactive_items(
    db_session: Session,
) -> None:
    body = _workshop(db_session, "Каховка", "Кузовной")
    mech = _workshop(db_session, "Каховка", "Слесарный")
    active = finance_service.create_cost_item(db_session, name="Аренда", is_active=True)
    old_used = finance_service.create_cost_item(db_session, name="Старая", is_active=True)
    finance_service.create_cost_item(db_session, name="Забытая", is_active=False)
    for workshop, item, amount in (
        (body, active, "100"),
        (mech, active, "7"),
        (body, old_used, "3"),
    ):
        finance_service.set_expense(
            db_session,
            cost_item_id=item.id,
            workshop_id=workshop.id,
            year=2026,
            month=9,
            amount=Decimal(amount),
        )
    old_used.is_active = False  # now inactive, but holds a value this year
    db_session.commit()

    result = finance_service.get_expenses_year(db_session, 2026)

    assert sorted(i.name for i in result["items"]) == ["Аренда", "Старая"]
    by_workshop = {w["workshop_id"]: w for w in result["workshops"]}
    assert set(by_workshop) == {body.id, mech.id}
    assert by_workshop[body.id]["totals"][8] == Decimal("103")
    assert by_workshop[mech.id]["totals"][8] == Decimal("7")
    assert all(len(row["months"]) == 12 for w in result["workshops"] for row in w["rows"])
    # another year is empty
    assert finance_service.get_expenses_year(db_session, 2025)["workshops"][0]["totals"][8] == 0


def test_api_cost_items_and_expenses_cycle(client, db_session: Session, auth_headers) -> None:
    workshop = _workshop(db_session)
    db_session.commit()

    created = client.post(
        ITEMS_URL, headers=auth_headers, json={"name": "  Аренда  ", "is_active": True}
    )
    assert created.status_code == 201
    assert created.json()["name"] == "Аренда"
    item_id = created.json()["id"]
    assert client.post(ITEMS_URL, headers=auth_headers, json={"name": "Аренда"}).status_code == 409
    assert client.post(ITEMS_URL, headers=auth_headers, json={"name": "   "}).status_code == 422

    saved = client.put(
        EXPENSES_URL,
        headers=auth_headers,
        json={
            "cost_item_id": item_id,
            "workshop_id": str(workshop.id),
            "year": 2026,
            "month": 9,
            "amount": "1500.50",
        },
    )
    assert saved.status_code == 204
    year = client.get(EXPENSES_URL, headers=auth_headers, params={"year": 2026}).json()
    assert year["workshops"][0]["rows"][0]["months"][8] == "1500.50"

    # deactivate -> the cell can no longer be edited, the amount stays
    client.put(
        f"{ITEMS_URL}/{item_id}", headers=auth_headers, json={"name": "Аренда", "is_active": False}
    )
    refused = client.put(
        EXPENSES_URL,
        headers=auth_headers,
        json={
            "cost_item_id": item_id,
            "workshop_id": str(workshop.id),
            "year": 2026,
            "month": 9,
            "amount": "1",
        },
    )
    assert refused.status_code == 422
    assert client.delete(f"{ITEMS_URL}/{item_id}", headers=auth_headers).status_code == 409
    missing = client.delete(f"{ITEMS_URL}/{uuid.uuid4()}", headers=auth_headers)
    assert missing.status_code == 404


def test_items_follow_the_hand_set_order_new_ones_go_last(db_session: Session) -> None:
    a = finance_service.create_cost_item(db_session, name="Аренда", is_active=True)
    b = finance_service.create_cost_item(db_session, name="Бензин", is_active=True)
    c = finance_service.create_cost_item(db_session, name="Вода", is_active=True)
    assert [i.name for i in finance_service.list_cost_items(db_session)] == [
        "Аренда",
        "Бензин",
        "Вода",
    ]

    finance_service.reorder_cost_items(db_session, [c.id, a.id, b.id])
    assert [i.name for i in finance_service.list_cost_items(db_session)] == [
        "Вода",
        "Аренда",
        "Бензин",
    ]

    finance_service.create_cost_item(
        db_session, name="Андреев", is_active=True
    )  # alphabetically first
    names = [i.name for i in finance_service.list_cost_items(db_session)]
    assert names == ["Вода", "Аренда", "Бензин", "Андреев"]  # new item is last

    # a partial list: the listed ones first, the rest keep their order after
    finance_service.reorder_cost_items(db_session, [b.id])
    names = [i.name for i in finance_service.list_cost_items(db_session)]
    assert names == ["Бензин", "Вода", "Аренда", "Андреев"]

    with pytest.raises(FinanceInvalid):
        finance_service.reorder_cost_items(db_session, [a.id, a.id])
    with pytest.raises(FinanceNotFound):
        finance_service.reorder_cost_items(db_session, [uuid.uuid4()])
    # and the Расходы table lists them in that order
    _workshop(db_session)
    db_session.commit()
    year = finance_service.get_expenses_year(db_session, 2026)
    assert [i.name for i in year["items"]] == names


def test_api_reorder(client, auth_headers) -> None:
    ids = [
        client.post(ITEMS_URL, headers=auth_headers, json={"name": n}).json()["id"]
        for n in ("Аренда", "Бензин")
    ]
    response = client.put(f"{ITEMS_URL}/order", headers=auth_headers, json={"ids": ids[::-1]})
    assert response.status_code == 200
    assert [i["name"] for i in response.json()] == ["Бензин", "Аренда"]
    assert [i["name"] for i in client.get(ITEMS_URL, headers=auth_headers).json()] == [
        "Бензин",
        "Аренда",
    ]


# ---- import of the P&L workbook ---------------------------------------------------


def _pnl_workbook(title: str = "Сентябрь 2026") -> bytes:
    """The layout of PnL/ЧП Новая.xlsm: workshops merged in row 1 (comment,
    Доход, Расход each), cost items from row 3, "V" rows = breakdown lines,
    "Итого" closes the table."""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["B1"] = "Статья доходов и расходов"
    ws.merge_cells("B1:C2")
    for col, label in (("D", "Каховка Кузов"), ("G", "Каховка Слесарка"), ("J", "Солнцево Кузов")):
        ws[f"{col}1"] = label
        start = ws[f"{col}1"].column
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + 2)
        ws.cell(row=2, column=start, value="В том числе")
        ws.cell(row=2, column=start + 1, value="Доход")
        ws.cell(row=2, column=start + 2, value="Расход")
    ws["B3"] = "Выручка (страховые)"
    ws["E3"] = 979371.7  # income only: brings nothing
    ws["B4"] = "Аренда "
    ws["F4"], ws["I4"], ws["L4"] = 626105, 591250.5, 0
    ws["B5"] = "Реклама"
    ws["F5"] = 100
    ws["I5"] = -5  # negative: skipped with a warning
    ws["B6"] = "Расходы по оборудованию"
    ws["F6"] = 122340
    ws["C7"], ws["D7"], ws["F7"] = "V", "в том числе: компрессор", 16225  # breakdown row
    ws["B8"] = "Новая статья"
    ws["I8"] = 7
    ws["B9"] = "Старая"
    ws["F9"] = 50
    ws["B10"] = "Итого"
    ws["F10"] = "=SUM(F4:F9)"
    ws["B11"] = "Расход после итога"
    ws["F11"] = 999
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _import_world(db: Session):
    body = _workshop(db, "Каховка", "Кузовной")
    mech = _workshop(db, "Каховка", "Слесарный")
    sol = _workshop(db, "Солнцево", "Кузовной")
    finance_service.create_cost_item(db, name="Аренда", is_active=True)
    finance_service.create_cost_item(db, name="Реклама", is_active=True)
    finance_service.create_cost_item(db, name="Расходы по оборудованию", is_active=True)
    finance_service.create_cost_item(db, name="Старая", is_active=False)
    return body, mech, sol


def test_import_preview_reads_the_layout_and_changes_nothing(db_session: Session) -> None:
    body, mech, sol = _import_world(db_session)

    report = finance_import_service.run_import(
        db_session, _pnl_workbook(), year=2026, month=9, apply=False
    )

    assert report["applied"] is False
    assert [g["workshop"] for g in report["workshops"]] == [
        "Каховка — Кузовной",
        "Каховка — Слесарный",
        "Солнцево — Кузовной",
    ]
    # Аренда x3 (the 0 counts), Реклама Каховка Кузов only (-5 skipped),
    # Расходы по оборудованию x1 ("V" row ignored); "Новая статья" is not in the
    # directory (not loaded, reported), the inactive "Старая" is skipped,
    # revenue/Итого/after-Итого rows bring nothing.
    assert report["cells"] == 5
    assert report["total"] == Decimal("626105") + Decimal("591250.50") + Decimal("100") + Decimal(
        "122340"
    )
    assert report["missing_items"] == ["Новая статья"]
    assert report["skipped_inactive"] == ["Старая"]
    assert any("отрицательное" in w for w in report["warnings"])
    assert db_session.query(CostExpense).count() == 0
    assert db_session.query(CostItem).count() == 4  # nothing is ever created


def test_import_loads_the_chosen_month_creates_nothing_and_replaces_values(
    db_session: Session,
) -> None:
    body, mech, sol = _import_world(db_session)
    rent = db_session.query(CostItem).filter_by(name="Аренда").one()
    finance_service.set_expense(
        db_session,
        cost_item_id=rent.id,
        workshop_id=body.id,
        year=2026,
        month=10,
        amount=Decimal("1"),
    )

    report = finance_import_service.run_import(
        db_session, _pnl_workbook(), year=2026, month=10, apply=True
    )

    assert report["applied"] is True
    # nothing is created: the item missing from the directory is only reported
    assert db_session.query(CostItem).filter_by(name="Новая статья").count() == 0
    assert report["missing_items"] == ["Новая статья"]
    year = finance_service.get_expenses_year(db_session, 2026)
    by_workshop = {w["workshop_id"]: w for w in year["workshops"]}
    # October, index 9: replaced, and the other months untouched
    assert by_workshop[body.id]["totals"][9] == Decimal("626105") + Decimal("100") + Decimal(
        "122340"
    )
    assert by_workshop[mech.id]["totals"][9] == Decimal("591250.50")
    assert by_workshop[sol.id]["totals"][9] == Decimal("0")
    assert by_workshop[body.id]["totals"][8] == Decimal("0")

    # importing again does not duplicate anything
    finance_import_service.run_import(db_session, _pnl_workbook(), year=2026, month=10, apply=True)
    assert db_session.query(CostItem).count() == 4
    assert db_session.query(CostExpense).count() == 4


def test_import_rejects_garbage_and_sheets_without_workshop_columns(db_session: Session) -> None:
    _import_world(db_session)
    with pytest.raises(FinanceInvalid):
        finance_import_service.run_import(
            db_session, b"not an excel file", year=2026, month=9, apply=False
        )

    wb = Workbook()
    wb.active["B1"] = "пусто"
    buffer = io.BytesIO()
    wb.save(buffer)
    with pytest.raises(FinanceInvalid):
        finance_import_service.run_import(
            db_session, buffer.getvalue(), year=2026, month=9, apply=False
        )


def test_api_import_dry_run_then_apply(client, db_session: Session, auth_headers) -> None:
    _import_world(db_session)
    db_session.commit()
    url = "/api/v1/finance/expenses/import"
    headers = {**auth_headers, "Content-Type": "application/octet-stream"}

    preview = client.post(
        url,
        headers=headers,
        params={"year": 2026, "month": 9, "dry_run": "true"},
        content=_pnl_workbook(),
    )
    assert preview.status_code == 200
    assert preview.json()["applied"] is False
    assert db_session.query(CostExpense).count() == 0

    loaded = client.post(
        url,
        headers=headers,
        params={"year": 2026, "month": 9, "dry_run": "false"},
        content=_pnl_workbook(),
    )
    assert loaded.status_code == 200
    assert loaded.json()["applied"] is True
    assert loaded.json()["missing_items"] == ["Новая статья"]
    assert db_session.query(CostExpense).count() == 4

    broken = client.post(url, headers=headers, params={"year": 2026, "month": 9}, content=b"junk")
    assert broken.status_code == 422
