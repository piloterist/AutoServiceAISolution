"""Касса -> "В Excel" (product ask, 2026-10-27): every operation of one cash desk
over a period as an Excel table - Год / Месяц / День / Операция / Заказ-наряд /
Комментарий / Сумма.

Every day of the period has at least two rows, "Остаток на начало" and
"Остаток на конец" (the cash balance, card payments never touch it), with the
day's operations between them - one row per cash entry, so its work order and
comment are visible. "Расход НАЛ" amounts are negative (they leave the drawer).
There are no gaps and the sheet filters/pivots cleanly.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.cash_entry import (
    KIND_CARD_IN,
    KIND_CASH_IN,
    KIND_CASH_OUT,
    KIND_LABELS,
    CashEntry,
)
from app.services import cash_service
from app.services.cash_service import CashAccess, CashInvalid

OPENING = "Остаток на начало"
CLOSING = "Остаток на конец"
# The order of the operations inside a day (between the two balances).
OPERATION_ORDER = (KIND_CASH_IN, KIND_CARD_IN, KIND_CASH_OUT)
# Nobody needs more; keeps one request from building a huge sheet.
MAX_DAYS = 3660

MONTHS = (
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
)

# (day, operation, work order number, comment, amount)
Row = tuple[date, str, str | None, str | None, Decimal]


def export_rows(
    db: Session,
    access: CashAccess,
    workshop_id: uuid.UUID,
    date_from: date,
    date_to: date,
) -> list[Row]:
    """One row per cash entry plus the two balance rows of every day of
    [date_from, date_to], oldest first: the opening balance, the day's entries
    (by kind, then entry time; an expense is negative), the closing balance
    (= the next day's opening)."""
    cash_service._require(access, workshop_id)
    if date_to < date_from:
        raise CashInvalid("Дата «по» не может быть раньше даты «с»")
    if (date_to - date_from).days + 1 > MAX_DAYS:
        raise CashInvalid("Слишком длинный период для выгрузки")

    before = CashEntry.entry_date < date_from
    balance = cash_service._sum(db, workshop_id, KIND_CASH_IN, before) - cash_service._sum(
        db, workshop_id, KIND_CASH_OUT, before
    )

    by_day: dict[date, list[CashEntry]] = {}
    for entry in db.scalars(
        select(CashEntry)
        .where(
            CashEntry.workshop_id == workshop_id,
            CashEntry.entry_date >= date_from,
            CashEntry.entry_date <= date_to,
        )
        .order_by(CashEntry.entry_date, CashEntry.created_at)
    ):
        by_day.setdefault(entry.entry_date, []).append(entry)

    rows: list[Row] = []
    day = date_from
    while day <= date_to:
        # stable sort: the entry time order is kept inside a kind
        entries = sorted(by_day.get(day, []), key=lambda e: OPERATION_ORDER.index(e.kind))
        rows.append((day, OPENING, None, None, cash_service._money(balance)))
        for entry in entries:
            sign = -1 if entry.kind == KIND_CASH_OUT else 1
            rows.append(
                (
                    day,
                    KIND_LABELS[entry.kind],
                    entry.work_order_number,
                    entry.comment,
                    cash_service._money(sign * entry.amount),
                )
            )
            if entry.kind == KIND_CASH_IN:
                balance += entry.amount
            elif entry.kind == KIND_CASH_OUT:
                balance -= entry.amount
        rows.append((day, CLOSING, None, None, cash_service._money(balance)))
        day += timedelta(days=1)
    return rows


def build_workbook(rows: list[Row]) -> bytes:
    """The Excel file: one formatted table with a filter on every column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Касса"

    headers = ["Год", "Месяц", "День", "Операция", "Заказ-наряд", "Комментарий", "Сумма"]
    ws.append(headers)
    for day, operation, work_order, comment, amount in rows:
        ws.append(
            [day.year, MONTHS[day.month - 1], day, operation, work_order, comment, float(amount)]
        )

    last_row = len(rows) + 1
    for row in range(2, last_row + 1):
        ws.cell(row=row, column=1).number_format = "0"
        ws.cell(row=row, column=3).number_format = "DD.MM.YYYY"
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="left")
        ws.cell(row=row, column=7).number_format = "#,##0.00"

    table = Table(displayName="KassaOperacii", ref=f"A1:G{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)  # a table carries the filter buttons on every header

    for col, width in zip(range(1, 8), (8, 14, 14, 20, 16, 40, 16), strict=True):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
